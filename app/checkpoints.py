"""Readable checkpoint exports; no model requests are made."""
import io
import re
import os
from pathlib import Path
from html import escape


def readable_math(expression):
    value=expression.strip()
    for _ in range(8):
        old=value
        value=re.sub(r'\\(?:d?frac)\{([^{}]*)\}\{([^{}]*)\}',lambda m:'('+m[1]+')/('+m[2]+')',value)
        value=re.sub(r'\\sqrt\{([^{}]*)\}',lambda m:'√('+m[1]+')',value)
        if value==old:break
    symbols={'times':'×','cdot':'·','div':'÷','pm':'±','leq':'≤','le':'≤','geq':'≥','ge':'≥','neq':'≠','infty':'∞','sum':'∑','int':'∫','pi':'π','theta':'θ','alpha':'α','beta':'β','gamma':'γ','Delta':'Δ','delta':'δ','rightarrow':'→','to':'→','approx':'≈','in':'∈'}
    value=re.sub(r'\\([A-Za-z]+)',lambda m:symbols.get(m[1],'' if m[1] in {'left','right','displaystyle'} else m[1]),value)
    supers=str.maketrans('0123456789+-=()n','⁰¹²³⁴⁵⁶⁷⁸⁹⁺⁻⁼⁽⁾ⁿ')
    subs=str.maketrans('0123456789+-=()','₀₁₂₃₄₅₆₇₈₉₊₋₌₍₎')
    value=re.sub(r'\^\{?([0-9n+-]+)\}?',lambda m:m[1].translate(supers),value)
    value=re.sub(r'_\{?([0-9+-]+)\}?',lambda m:m[1].translate(subs),value)
    value=re.sub(r'(?:text|mathrm|mathbf|operatorname)\{([^{}]*)\}',r'\1',value)
    value=value.replace(r'\,',' ').replace(r'\;',' ').replace(r'\!','').replace(r'\%','%')
    return value


def readable_text(text):
    value=str(text)
    code=[]
    def protect(match):
        code.append(match[1] if match[1] is not None else match[2]);return '\x00CODE'+str(len(code)-1)+'\x00'
    value=re.sub(r'```[^\n]*\n([\s\S]*?)```|`([^`\n]+)`',protect,value)
    value=value.translate(str.maketrans({'\u2010':'-','\u2011':'-','\u202f':' ','\u00a0':' ','\u00ad':'','\u200b':''}))
    value=re.sub(r'\{\s*:?(?:reasoning|execution|relevance|analysis|final|outcome|action)\s*\}','',value,flags=re.I)
    value=re.sub(r'\$\$([\s\S]+?)\$\$|\\\[([\s\S]+?)\\\]|\\\((.+?)\\\)|\$(?!\s)([^$\n]+?)\$',lambda m:readable_math(next(g for g in m.groups() if g is not None)),value)
    value=re.sub(r'(?<!\w)\*\*(?=\S)(.+?)(?<=\S)\*\*(?!\w)',r'\1',value)
    value=re.sub(r'(?<!\w)\*(?=\S)([^*\n]+?)(?<=\S)\*(?!\w)',r'\1',value)
    value=re.sub(r'\[([^\]]+)\]\(([^)]+)\)',r'\1 (\2)',value)
    value=re.sub(r'(?m)^#{1,6}\s+', '',value)
    value=re.sub(r'\x00CODE(\d+)\x00',lambda m:code[int(m[1])],value)
    return value.strip()


def sections(state):
    plan=state.get('curriculum') or {}
    objectives=plan.get('objectives',[])
    modules=plan.get('modules',[])
    idx=state.get('objective_index',0)
    current_obj=objectives[min(idx,len(objectives)-1)] if objectives else {}
    yield 'Personal Learning checkpoint',plan.get('title','Learning session')
    yield 'Goal',state.get('profile',{}).get('goal','')
    pos_parts=[f"Stage: {state.get('stage','')}"]
    if modules and current_obj.get('module_title'):
        pos_parts.append(f"Module: {current_obj['module_title']}")
    if objectives:
        pos_parts.append(f"Objective {min(idx+1,len(objectives))} of {len(objectives)} ({current_obj.get('title','')})")
    next_act=(state.get('next_action') or {}).get('label','')
    if next_act:
        pos_parts.append(f"Next step: {next_act}")
    yield 'Progress',' · '.join(pos_parts)
    if plan.get('intended_outcome') or plan.get('scope_note') or plan.get('grounding_note'):
        scope_lines=[x for x in (plan.get('intended_outcome',''),plan.get('scope_note',''),plan.get('grounding_note','')) if x]
        if scope_lines:
            yield 'Scope and grounding','\n'.join(scope_lines)
    if modules:
        mod_lines=[]
        obj_by_id={o['id']:o for o in objectives if isinstance(o,dict) and o.get('id')}
        for m_idx,mod in enumerate(modules,1):
            mod_lines.append(f"Module {m_idx}: {mod.get('title','')}" + (f" — {mod['summary']}" if mod.get('summary') else ''))
            for oid in mod.get('objective_ids',[]):
                obj=obj_by_id.get(oid)
                if obj:
                    concepts=', '.join(c.get('title','') for c in obj.get('concepts',[]) if isinstance(c,dict) and c.get('title'))
                    mod_lines.append(f"  - {obj['title']} ({obj.get('minutes',10)} min): {obj.get('outcome','')}" + (f" [Concepts: {concepts}]" if concepts else ''))
        if mod_lines:
            yield 'Curriculum hierarchy','\n'.join(mod_lines)
    for i,item in enumerate(objectives,1):
        yield f"{i}. {item['title']}",item.get('outcome','')
    if plan.get('optional_enrichment'):
        yield 'Optional enrichment','\n'.join(f"- {x}" for x in plan['optional_enrichment'] if x)
    if state.get('lesson'):
        yield 'Current lesson',state['lesson']['explanation']
        yield 'Worked example','\n'.join(f'{i}. {step}' for i,step in enumerate(state['lesson']['worked_example'],1))
    current=state.get('lesson')
    seen={str(current)} if current else set()
    for entry in state.get('lesson_history',[]):
        lesson=entry.get('lesson')
        if not lesson or str(lesson) in seen:continue
        seen.add(str(lesson))
        yield 'Saved lesson: '+entry['title'],lesson.get('explanation','')
        if lesson.get('worked_example'):
            yield 'Worked example','\n'.join(f'{i}. {step}' for i,step in enumerate(lesson['worked_example'],1))
    if state.get('repair'):
        yield 'Repair',state['repair']['explanation']
    review_map=state.get('review') or {}
    if review_map:
        rev_lines=[]
        for cid,rec in review_map.items():
            if isinstance(rec,dict):
                rev_lines.append(f"{cid}: status={rec.get('status','unassessed')}, taught={rec.get('taught',False)}, next_due={rec.get('next_due','')}, evidence={rec.get('evidence','')}")
        if rev_lines:
            yield 'Concept and review state','\n'.join(rev_lines)
    source=state.get('source',{})
    yield 'Reference',source.get('title','')+'\n'+source.get('origin','')


def export(state,format):
    content=[(readable_text(title),readable_text(text)) for title,text in sections(state)]
    if format=='md':
        return ('\n\n'.join(('# ' if i==0 else '## ')+title+'\n\n'+text for i,(title,text) in enumerate(content))+'\n').encode('utf-8'),'text/markdown'
    if format=='docx':
        from docx import Document
        from docx.shared import Pt
        doc=Document();doc.styles['Normal'].font.name='Arial';doc.styles['Normal'].font.size=Pt(11)
        for i,(title,text) in enumerate(content):
            doc.add_heading(title,0 if i==0 else 1)
            doc.add_paragraph(text)
        output=io.BytesIO();doc.save(output)
        return output.getvalue(),'application/vnd.openxmlformats-officedocument.wordprocessingml.document'
    if format=='pdf':
        from reportlab.platypus import SimpleDocTemplate,Paragraph,Spacer
        from reportlab.lib.styles import getSampleStyleSheet
        from reportlab.pdfbase import pdfmetrics
        from reportlab.pdfbase.ttfonts import TTFont
        candidates=[os.getenv('CHECKPOINT_PDF_FONT',''),r'C:\Windows\Fonts\msyh.ttc',r'C:\Windows\Fonts\arial.ttf','/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf']
        font='Helvetica'
        for candidate in candidates:
            if candidate and Path(candidate).is_file():
                try:
                    pdfmetrics.registerFont(TTFont('CheckpointText',candidate))
                    font='CheckpointText'
                    break
                except Exception:
                    continue
        styles=getSampleStyleSheet()
        for style in styles.byName.values():style.fontName=font
        output=io.BytesIO();story=[]
        for i,(title,text) in enumerate(content):
            story.extend([Paragraph(escape(title),styles['Title' if i==0 else 'Heading2']),Paragraph(escape(text).replace('\n','<br/>'),styles['BodyText']),Spacer(1,12)])
        SimpleDocTemplate(output,title='Personal Learning checkpoint',leftMargin=48,rightMargin=48,topMargin=48,bottomMargin=48).build(story)
        return output.getvalue(),'application/pdf'
    raise ValueError('Choose Markdown, DOCX or PDF.')
