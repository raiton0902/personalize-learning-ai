import io
import json
import socket
from copy import deepcopy

import pytest
from docx import Document
from starlette.testclient import TestClient

import web_server
import web_store
import web_tutor
import references


@pytest.fixture
def client(monkeypatch, tmp_path):
    monkeypatch.setenv("WEB_DB", str(tmp_path / "web.sqlite3"))
    monkeypatch.setenv("TRANSCRIPTION_ALLOW_PAID", "false")
    monkeypatch.setattr(web_tutor, "propose", lambda state, changes="": ({"title":"Python basics", "scope_note":"A small course", "objectives":[{"id":"python.variables","title":"Variables","outcome":"Assign and inspect a value", "prerequisites":[],"source_ids":[state["source"]["id"]],"minutes":10}]}, {}))
    monkeypatch.setattr(web_tutor, "teach", lambda state: ({"explanation":"A variable refers to a value.","worked_example":["Set x to 3.","Print x."]}, {}))
    monkeypatch.setattr(web_tutor, "practice", lambda *args: ({"id":"generated","prompt":"What is x after x=3?","answer":"3","solution":"The assignment sets x to 3."}, {}))
    monkeypatch.setattr(web_tutor, "assess", lambda *args: ({"status":"incorrect","evidence":"The answer differs.","possible_gap":"assignment","diagnostic_question":"What does the assignment do?"}, {}))
    monkeypatch.setattr(web_tutor, "repair", lambda *args: ({"explanation":"Assignment binds the name to a value.","observation":"Needed help with assignment."}, {}))
    with TestClient(web_server.app) as value:
        # These downstream workflow cases explicitly opt out of intake; new intake tests use a separate client.
        post = value.post
        def scoped_post(url, **kwargs):
            if url == '/api/courses' and 'json' in kwargs:
                kwargs['json'] = {'skip_intake': True, **kwargs['json']}
            return post(url, **kwargs)
        value.post = scoped_post
        yield value


def signup(client, email="learner@example.test"):
    response = client.post("/api/auth/signup", json={"email":email,"password":"a-long-test-password"})
    assert response.status_code == 200, response.text
    assert "httponly" in response.headers["set-cookie"].lower()
    return response


def test_home_metadata_and_local_favicon(client):
    from bs4 import BeautifulSoup
    response=client.get('/');assert response.status_code==200
    page=BeautifulSoup(response.text,'html.parser')
    assert page.find('meta',attrs={'name':'description'})['content']
    assert page.find('link',rel='canonical')['href']=='/'
    assert page.find('meta',property='og:title')['content']=='Personal Learning'
    assert page.find('meta',property='og:description')['content']
    assert page.find('meta',attrs={'name':'twitter:card'})['content']=='summary'
    assert page.find('meta',property='og:image') is None
    icon=page.find('link',rel='icon')['href']
    assert client.get(icon).status_code==200


def test_missing_page_has_heading_and_home_link_while_api_stays_json(client):
    from bs4 import BeautifulSoup
    response=client.get('/missing-page');assert response.status_code==404
    page=BeautifulSoup(response.text,'html.parser')
    assert page.h1.get_text()=='Page not found' and page.find('a',href='/')
    response=client.get('/api/missing-page')
    assert response.status_code==404 and response.json()['code']=='not_found'


def test_accounts_profile_ownership_and_logout(client):
    assert client.get("/api/me").json()["user"] is None
    assert client.get("/api/resources").status_code == 401
    signup(client)
    profile={"language":"Mandarin","level":"I get stuck with loops","tone":"Short steps"}
    assert client.post("/api/profile",json=profile).status_code == 200
    assert 'language' not in client.get('/api/me').json()['user']['profile']
    assert client.get('/api/me').json()['user']['profile']['tone']=='Short steps'
    resource=client.post("/api/resources",json={"text":"An original explanation of assigning Python variables.","title":"Variables"}).json()["resource"]
    assert resource["id"] != "pasted"
    client.post("/api/logout")
    assert client.get("/api/resources").status_code == 401
    signup(client,"other@example.test")
    assert client.get("/api/resources").json()["resources"] == []
    response=client.post("/api/courses",json={"profile":{"goal":"learn Python"},"resource_id":resource["id"]})
    assert response.status_code == 400
    client.post("/api/logout")
    assert client.post("/api/auth/login",json={"email":"learner@example.test","password":"wrong"}).status_code == 400
    assert client.post("/api/auth/login",json={"email":"learner@example.test","password":"a-long-test-password"}).status_code == 200


def test_real_upload_plan_approval_diagnostic_resume_and_export(client):
    signup(client)
    doc=Document();doc.add_paragraph("Python variables refer to values. Assign x = 3 and inspect x.")
    data=io.BytesIO();doc.save(data)
    response=client.post("/api/resources",files={"file":("source.docx",data.getvalue(),"application/vnd.openxmlformats-officedocument.wordprocessingml.document")})
    assert response.status_code == 200,response.text
    rid=response.json()["resource"]["id"]
    response=client.post("/api/courses",json={"profile":{"goal":"Learn Python variables","language":"automatic"},"resource_id":rid})
    assert response.status_code == 200,response.text
    cid=response.json()["id"]
    assert client.post(f"/api/courses/{cid}",json={"action":"teach"}).status_code == 400
    for action in ["approve","teach","practice"]:
        response=client.post(f"/api/courses/{cid}",json={"action":action})
        assert response.status_code == 200,response.text
    assert "answer" not in response.json()["state"]["question"]
    response=client.post(f"/api/courses/{cid}",json={"action":"attempt","answer":"4","work":"I added one."})
    assert response.json()["state"]["stage"] == "diagnostic"
    response=client.post(f"/api/courses/{cid}",json={"action":"diagnostic","text":"I thought assignment adds one."})
    assert response.json()["state"]["stage"] == "repair"
    assert len(client.get("/api/activity").json()["events"]) >= 5
    assert "Python basics" in client.get(f"/api/courses/{cid}/export").text
    client.post("/api/logout")
    client.post("/api/auth/login",json={"email":"learner@example.test","password":"a-long-test-password"})
    assert client.get(f"/api/courses/{cid}").json()["state"]["stage"] == "repair"


def test_private_url_and_video_setup_are_explicit(client):
    signup(client)
    for url in ["http://127.0.0.1/private","http://169.254.169.254/latest/meta-data/","file:///etc/passwd"]:
        response=client.post("/api/resources",json={"kind":"url","url":url})
        assert response.status_code == 400
    response=client.post("/api/resources",files={"file":("clip.mp4",b"fake video","video/mp4")})
    assert response.status_code == 400
    assert "transcription is not configured" in response.json()["error"]
    response=client.post("/api/resources",files={"file":("captions.vtt",b"WEBVTT\n\n00:00.000 --> 00:05.000\nVariables refer to values; x equals three.","text/vtt")})
    assert response.status_code == 200


def test_public_html_is_read_with_pinned_address(monkeypatch):
    monkeypatch.setattr(socket,"getaddrinfo",lambda *args,**kwargs:[(socket.AF_INET,socket.SOCK_STREAM,6,"",("93.184.216.34",443))])
    class FakeResponse:
        status=200
        def getheader(self,key,default=""): return "text/html"
        def read(self,n): return b"<title>Course</title><nav>ignore</nav><main>Python variables refer to values, and assignment binds a value to a name.</main>"
    class FakeConnection:
        def __init__(self,hostname,address,port): assert address=="93.184.216.34"
        def request(self,*args,**kwargs): pass
        def getresponse(self): return FakeResponse()
        def close(self): pass
    monkeypatch.setattr(references,"PinnedHTTPS",FakeConnection)
    result=references.read_url("https://example.org/course")
    assert result["title"] == "Course"
    assert "ignore" not in result["text"]


def test_cross_origin_post_rejected(client):
    response=client.post("/api/auth/signup",json={"email":"a@example.test","password":"test-password-123"},headers={"Origin":"https://other.example"})
    assert response.status_code == 403


def test_configured_audio_extraction_uses_service_transport_without_live_call(monkeypatch):
    import shutil
    import wave
    if not shutil.which("ffmpeg"):
        pytest.skip("FFmpeg not installed")
    monkeypatch.setenv("TRANSCRIPTION_BASE_URL","https://speech.example/v1")
    monkeypatch.setenv("TRANSCRIPTION_MODEL","test-pretrained-model")
    monkeypatch.setenv("TRANSCRIPTION_API_KEY","synthetic-test-key")
    monkeypatch.setenv("TRANSCRIPTION_ALLOW_PAID","true")
    captured=[]
    class Response:
        status_code=200
        def json(self):return {"text":"Python assignment binds a name to a value. This is a synthetic transcript."}
    def post(url,**kwargs):
        captured.append(url)
        assert kwargs["files"]["file"][2] == "audio/wav"
        return Response()
    monkeypatch.setattr(references.requests,"post",post)
    audio=io.BytesIO()
    with wave.open(audio,"wb") as wav:
        wav.setnchannels(1);wav.setsampwidth(2);wav.setframerate(16000);wav.writeframes(b"\0\0"*16000)
    result=references.document(audio.getvalue(),"example.wav")
    assert result["id"] == "transcript"
    assert captured == ["https://speech.example/v1/audio/transcriptions"]


def test_redirect_to_private_address_is_rejected(monkeypatch):
    original=references.public_target
    calls=[]
    def target(url):
        calls.append(url)
        if url.startswith("https://public.example"):
            from urllib.parse import urlsplit
            return urlsplit(url),443,"93.184.216.34"
        return original(url)
    class Redirect:
        status=302
        def getheader(self,key): return "http://127.0.0.1/private"
    class Connection:
        def __init__(self,*args):pass
        def request(self,*args,**kwargs):pass
        def getresponse(self):return Redirect()
        def close(self):pass
    monkeypatch.setattr(references,"public_target",target)
    monkeypatch.setattr(references,"PinnedHTTPS",Connection)
    with pytest.raises(ValueError,match="Private"):
        references.read_url("https://public.example/start")
    assert len(calls)==2


def test_general_plan_rejects_untrusted_source_ids(monkeypatch):
    state={"source":{"id":"reference","text":"some source"},"profile":{"goal":"learn Python"},"curriculum":None,"objective_index":0}
    monkeypatch.setattr(web_tutor,"generate",lambda *args:({"objectives":[{"id":"a","prerequisites":[],"source_ids":["unknown"]}]},{}))
    # Call the original implementation, not the route fixture's replacement.
    with pytest.raises(ValueError,match="references"):
        web_tutor.propose(state)


def test_saved_question_threads_preserve_stage_reuse_answers_and_keep_signals_private(client,monkeypatch):
    signup(client)
    result=client.post('/api/courses',json={'profile':{'goal':'Learn Python assignment'}}).json()
    cid=result['id']
    assert client.post(f'/api/courses/{cid}',json={'action':'question','text':'Why assignment?','request_id':'initial'}).status_code==400
    client.post(f'/api/courses/{cid}',json={'action':'approve'})
    client.post(f'/api/courses/{cid}',json={'action':'teach'})
    calls=[]
    def reply(state,question,thread):
        calls.append(question)
        return {'related':True,'new_goal':'','answer':'Assignment binds the name to its value.'},{}
    monkeypatch.setattr(web_tutor,'chat',reply)
    for _ in range(2):
        response=client.post(f'/api/courses/{cid}',json={'action':'question','text':'Why assignment?','request_id':'initial'})
        assert response.status_code==200,response.text
    state=response.json()['state']
    assert state['stage']=='lesson'
    assert len(state['followups'])==1
    assert len(calls)==1
    assert 'learning_signals' not in state
    account=web_store.account(client.cookies.get('learning_login'))
    stored,_=web_store.course(account['id'],cid)
    assert stored['learning_signals'][0]['kind']=='requested_help'
    assert 'Why assignment?' in stored['learning_signals'][0]['evidence']
    assert client.get(f'/api/courses/{cid}').json()['state']['followups'][0]['answer'].startswith('Assignment')


def test_question_evidence_is_passed_to_practice_without_fixed_ability_labels(monkeypatch):
    captured={}
    def generate(schema,purpose,data,instructions):
        captured.update(data=data,instructions=instructions)
        return {'prompt':'Assign x','answer':'x=3','solution':'Use assignment'},{}
    monkeypatch.setattr(web_tutor.tutor,'generate',generate)
    state={'source':{'id':'general','text':'Assignment binds a name.'},'profile':{'goal':'Python'},'objective_index':0,'curriculum':{'objectives':[{'id':'assignment'}]},'learning_signals':[{'objective_id':'assignment','evidence':'Asked: Why is assignment not addition?'}]}
    web_tutor.practice(state)
    assert captured['data']['recent_help_requests'][0]['evidence'].startswith('Asked:')
    assert 'not proof of inability' in captured['instructions']


def test_checkpoint_formats_are_real_files_and_private(client):
    from pypdf import PdfReader
    signup(client)
    result=client.post('/api/courses',json={'profile':{'goal':'Learn Python variables'}}).json()
    cid=result['id']
    pdf=client.get(f'/api/courses/{cid}/export?format=pdf')
    assert pdf.status_code==200
    assert pdf.content.startswith(b'%PDF-')
    assert len(PdfReader(io.BytesIO(pdf.content)).pages)>=1
    docx=client.get(f'/api/courses/{cid}/export?format=docx')
    assert docx.status_code==200
    assert docx.content.startswith(b'PK')
    assert 'Python basics' in '\n'.join(p.text for p in Document(io.BytesIO(docx.content)).paragraphs)
    assert client.get(f'/api/courses/{cid}/export?format=exe').status_code==400
    client.post('/api/logout')
    assert client.get(f'/api/courses/{cid}/export?format=pdf').status_code==401


def test_pattern_save_updates_account_and_course_atomically(client):
    signup(client)
    cid=client.post('/api/courses',json={'profile':{'goal':'Python'}}).json()['id']
    response=client.post('/api/profile',json={'goal':'Python','tone':'Friendly','course_id':cid})
    assert response.status_code==200
    assert client.get(f'/api/courses/{cid}').json()['state']['profile']['tone']=='Friendly'
    response=client.post('/api/profile',json={'goal':'Python','tone':'Should not persist','course_id':'not-owned'})
    assert response.status_code==400
    assert client.get('/api/me').json()['user']['profile']['tone']=='Friendly'


def test_universal_preferences_apply_across_courses_without_local_context_leak(client):
    signup(client)
    a=client.post('/api/courses',json={'profile':{'goal':'Algebra','level':'Stuck on equations','language':'Mandarin','context':'Exam A'}}).json()['id']
    b=client.post('/api/courses',json={'profile':{'goal':'Python','level':'','language':'automatic','context':''}}).json()['id']
    client.post('/api/profile',json={'goal':'Algebra','tone':'Friendly','interest':'Music','level':'Stuck on equations','language':'Mandarin','context':'Exam A','course_id':a})
    first=client.get(f'/api/courses/{a}').json()['state']['profile']
    second=client.get(f'/api/courses/{b}').json()['state']['profile']
    assert first['context']=='Exam A'
    assert second['context']=='' and second['level']=='' and second['language']=='automatic'
    assert second['tone']=='Friendly' and second['interest']=='Music'
    assert set(client.get('/api/me').json()['user']['profile'])=={'tone','interest'}


def test_chat_continuity_scope_decisions_and_account_isolation(client, monkeypatch):
    signup(client)
    cid=client.post('/api/courses',json={'profile':{'goal':'Python','context':'Old exam','level':'Old struggle'}}).json()['id']
    path=f'/api/courses/{cid}'
    client.post(path,json={'action':'approve'})
    calls=[]
    def chat(state,text,thread):
        calls.append(text)
        return {'related':text!='Plants?', 'new_goal':'Learn photosynthesis' if text=='Plants?' else '', 'answer':'Small concrete example.'},{}
    monkeypatch.setattr(web_tutor,'chat',chat)
    def send(action,**kwargs):
        response=client.post(path,json={'action':action,**kwargs})
        assert response.status_code==200,response.text
        return response.json()
    initial=send('question',text='Why?',request_id='one')
    tid=initial['thread_id']
    assert initial['state']['stage']=='ready'
    second=send('chat',thread_id=tid,text='Show an example',request_id='two')
    assert len(second['state']['followups'][0]['messages'])==4
    send('chat',thread_id=tid,text='Show an example',request_id='two')
    assert calls==['Why?','Show an example']
    pending=send('chat',thread_id=tid,text='Plants?',request_id='three')['state']['followups'][0]['pending']
    assert 'answer' not in pending
    discarded=send('chat_discard',thread_id=tid,pending_id=pending['id'])
    assert 'pending' not in discarded['state']['followups'][0]
    assert len(discarded['state']['followups'][0]['messages'])==4
    pending=send('chat',thread_id=tid,text='Plants?',request_id='four')['state']['followups'][0]['pending']
    # A failed plan leaves the pending decision intact and creates no course.
    original=web_tutor.propose
    def fail(state): raise ValueError('Temporary model failure')
    monkeypatch.setattr(web_tutor,'propose',fail)
    assert client.post(path,json={'action':'chat_transfer','thread_id':tid,'pending_id':pending['id']}).status_code==400
    assert len(client.get('/api/courses').json()['courses'])==1
    monkeypatch.setattr(web_tutor,'propose',original)
    transferred=send('chat_transfer',thread_id=tid,pending_id=pending['id'])
    new=transferred['state'];new_id=transferred['id']
    assert new_id!=cid and new['stage']=='plan' and new['approved_version'] is None
    assert new['profile']['context']=='' and new['profile']['level']=='' and new['resource_id'] is None
    assert new['profile']['goal']=='Learn photosynthesis'
    assert new['followups'][0]['messages'][0]['content']=='Plants?'
    assert client.post(f'/api/courses/{new_id}',json={'action':'teach'}).status_code==400
    assert client.post(path,json={'action':'chat_transfer','thread_id':tid,'pending_id':pending['id']}).status_code==400
    assert len(client.get('/api/courses').json()['courses'])==2
    assert 'pending' not in client.get(path).json()['state']['followups'][0]
    client.post('/api/logout');signup(client,'intruder@example.test')
    assert client.post(path,json={'action':'chat','thread_id':tid,'text':'Hi','request_id':'attack'}).status_code==400


def test_chat_prompt_contains_only_six_messages_from_this_thread(monkeypatch):
    captured={}
    def generate(schema,purpose,data,instruction):
        captured.update(data=data,instruction=instruction)
        return {'scope':'same_course','new_goal':'','answer':'Example'},{}
    monkeypatch.setattr(web_tutor.tutor,'generate',generate)
    state={'source':{'id':'general','text':'Python'},'profile':{'goal':'Python'},'objective_index':0,
           'curriculum':{'objectives':[{'id':'variables'}]},'lesson':{'explanation':'Assignment'},
           'learning_signals':[{'objective_id':'variables','evidence':'PRIVATE OTHER CHAT'}],
           'followups':[{'question':'OTHER CHAT SECRET'}]}
    thread={'objective_id':'variables','messages':[{'role':'user' if i%2==0 else 'assistant','content':str(i)} for i in range(10)]}
    web_tutor.chat(state,'Why?',thread)
    assert [m['content'] for m in captured['data']['recent_messages']]==list(map(str,range(4,10)))
    assert 'OTHER CHAT' not in json.dumps(captured['data'])
    assert 'explain WHY' in captured['instruction'] and 'concrete example' in captured['instruction']


def test_legacy_chat_is_migrated_and_new_off_topic_question_can_be_discarded(client,monkeypatch):
    signup(client)
    cid=client.post('/api/courses',json={'profile':{'goal':'Python'}}).json()['id']
    owner=web_store.account(client.cookies.get('learning_login'))['id']
    state,revision=web_store.course(owner,cid)
    state.update(approved_version=1,stage='lesson',followups=[{'id':'old','question':'Original','answer':'Saved answer','objective_id':'python.variables','objective_title':'Variables'}])
    web_store.save_course(owner,cid,state,revision)
    monkeypatch.setattr(web_tutor,'chat',lambda *args:({'related':True,'new_goal':'','answer':'Next answer'},{}))
    path=f'/api/courses/{cid}'
    result=client.post(path,json={'action':'chat','thread_id':'old','text':'Next','request_id':'next'}).json()
    assert [m['content'] for m in result['state']['followups'][0]['messages']]==['Original','Saved answer','Next','Next answer']
    monkeypatch.setattr(web_tutor,'chat',lambda *args:({'related':False,'new_goal':'Plants','answer':'Hidden'},{}))
    result=client.post(path,json={'action':'question','text':'Plants?','request_id':'new'}).json()
    tid=result['thread_id'];pending=result['state']['followups'][-1]['pending']['id']
    result=client.post(path,json={'action':'chat_discard','thread_id':tid,'pending_id':pending}).json()
    assert result['thread_id'] is None and len(result['state']['followups'])==1


def test_step_request_has_structured_steps_and_large_context_stays_bounded(monkeypatch):
    captured={}
    def generate(schema,purpose,data,instructions):
        captured.update(schema=schema,data=data,instructions=instructions)
        value={'scope':'same_course','new_goal':'','answer':'Assignment stores a value.',
               'steps':['Read the value 3 on the right.', 'Bind x to 3 so later uses of x refer to 3.']}
        return schema.model_validate(value).model_dump(),{}
    monkeypatch.setattr(web_tutor.tutor,'generate',generate)
    state={'source':{'id':'reference','text':'文'*10000},'profile':{'goal':'Python'},'objective_index':0,'curriculum':{'objectives':[{'id':'variables'}]}}
    thread={'objective_id':'variables','messages':[{'role':'assistant','content':'文'*2400} for _ in range(6)]}
    result,_=web_tutor.chat(state,'Explain x = 3 step by step',thread)
    assert '\n1. Read' in result['answer'] and '\n2. Bind' in result['answer']
    assert len(json.dumps(captured['data'],ensure_ascii=False).encode())<18000
    assert state['source']['text']=='文'*10000 and len(thread['messages'])==6


def test_checkpoints_omit_question_conversations_and_pending_topics():
    import checkpoints
    state={'profile':{'goal':'Python'},'curriculum':{'title':'Python','objectives':[]},
           'followups':[{'question':'Why?','answer':'First answer','messages':[{'role':'user','content':'Why?'},{'role':'assistant','content':'First answer'},{'role':'user','content':'Next?'},{'role':'assistant','content':'Second answer'}],
                         'pending':{'text':'UNCONFIRMED TOPIC','answer':'HIDDEN'}}]}
    text=str(list(checkpoints.sections(state)))
    assert 'Second answer' not in text and 'UNCONFIRMED' not in text and 'Why?' not in text


def test_similar_questions_offer_confirmation_without_another_model_call(client,monkeypatch):
    signup(client)
    cid=client.post('/api/courses',json={'profile':{'goal':'Learn algebra'}}).json()['id']
    path=f'/api/courses/{cid}';client.post(path,json={'action':'approve'})
    calls=[]
    def reply(*args):
        calls.append(args[1]);return {'related':True,'new_goal':'','answer':'Symbols can feel unfamiliar.'},{}
    monkeypatch.setattr(web_tutor,'chat',reply)
    first=client.post(path,json={'action':'question','text':'Why is it hard','request_id':'a'}).json()
    for text in ['Why is it hard?', 'Why this subject hard??']:
        response=client.post(path,json={'action':'question','text':text,'request_id':'b'}).json()
        assert response['similar_question']['id']==first['thread_id']
        assert len(response['state']['followups'])==1
    assert len(calls)==1
    separate=client.post(path,json={'action':'question','text':'Why this subject hard??','request_id':'c','similar_choice':'new'}).json()
    assert len(calls)==2 and len(separate['state']['followups'])==2
    other=client.post('/api/courses',json={'profile':{'goal':'Other course'}}).json()['id']
    client.post(f'/api/courses/{other}',json={'action':'approve'})
    response=client.post(f'/api/courses/{other}',json={'action':'question','text':'Why is it hard?','request_id':'d'}).json()
    assert 'similar_question' not in response and len(calls)==3


def test_unclear_messages_and_transport_failures_preserve_chat(client,monkeypatch):
    from llm import ModelRequestError
    signup(client)
    cid=client.post('/api/courses',json={'profile':{'goal':'Algebra'}}).json()['id']
    path=f'/api/courses/{cid}';client.post(path,json={'action':'approve'})
    calls=[]
    def fail(*args): calls.append(1);raise ModelRequestError('The request timed out.')
    monkeypatch.setattr(web_tutor,'chat',fail)
    unclear=client.post(path,json={'action':'question','text':'wdoiqndnioqwd','request_id':'a'}).json()
    assert 'wdoiqndnioqwd' in unclear['clarification'] and calls==[]
    timeout=client.post(path,json={'action':'question','text':'Explain variables','request_id':'b'}).json()
    assert 'Explain variables' in timeout['service_error'] and len(calls)==1
    assert 'followups' not in client.get(path).json()['state']
    monkeypatch.setattr(web_tutor,'chat',lambda *a:({'related':False,'clarify':True,'answer':'What do you mean?','new_goal':''},{}))
    unclear=client.post(path,json={'action':'question','text':'that thing somewhere','request_id':'c'}).json()
    assert 'that thing somewhere' in unclear['clarification']
    assert 'followups' not in client.get(path).json()['state']


def test_matching_keeps_math_negation_and_unfamiliar_language_distinct():
    from question_matching import similar_question,unclear_input
    state={'followups':[{'id':'one','objective_id':'o','question':'Why do we add 3 to x?','answer':'Answer'}]}
    assert similar_question(state,'Why do we add 4 to x?', 'o') is None
    assert similar_question(state,'Why do we not add 3 to x?', 'o') is None
    for text in ['photosynthesis','electromagnetism','解释一下变量','x = 3','Why?']:
        assert not unclear_input(text)


def test_generic_subject_difficulty_stays_in_course_even_if_model_misclassifies(monkeypatch):
    monkeypatch.setattr(web_tutor.tutor,'generate',lambda *a:({'scope':'new_learning','answer':'This subject uses unfamiliar symbols.','new_goal':'Some other subject'},{}))
    state={'source':{'id':'general','text':'Algebra'},'profile':{'goal':'Algebra'},'objective_index':0,'curriculum':{'objectives':[{'id':'o'}]}}
    result,_=web_tutor.chat(state,'Why this subject hard??',{'objective_id':'o','messages':[]})
    assert result['related'] and not result['clarify']


def test_model_clarification_is_not_mistaken_for_new_learning(monkeypatch):
    monkeypatch.setattr(web_tutor.tutor,'generate',lambda *a:({'scope':'new_learning','answer':'This phrase is unclear. Could you specify the concept?','new_goal':'Learn algebra concepts'},{}))
    state={'source':{'id':'general','text':'Algebra'},'profile':{'goal':'Algebra'},'objective_index':0,'curriculum':{'objectives':[{'id':'o'}]}}
    result,_=web_tutor.chat(state,'that thing somewhere maybe',{'objective_id':'o','messages':[]})
    assert result['clarify']


def test_old_false_topic_decision_can_be_answered_in_place(client,monkeypatch):
    signup(client)
    cid=client.post('/api/courses',json={'profile':{'goal':'Algebra'}}).json()['id']
    owner=web_store.account(client.cookies.get('learning_login'))['id']
    state,revision=web_store.course(owner,cid)
    state.update(approved_version=1,stage='lesson',followups=[{'id':'old','question':'Why this subject hard??','answer':'','objective_id':'python.variables','objective_title':'Variables','messages':[],
        'pending':{'id':'old-decision','text':'Why this subject hard??','goal':'Wrong new topic','answer':'Hidden'}}])
    web_store.save_course(owner,cid,state,revision)
    path=f'/api/courses/{cid}'
    assert client.get(path).json()['state']['followups'][0]['pending']['can_reconsider']
    monkeypatch.setattr(web_tutor,'chat',lambda *a:({'related':True,'answer':'Symbols are unfamiliar.','new_goal':''},{}))
    result=client.post(path,json={'action':'chat_reconsider','thread_id':'old','pending_id':'old-decision'}).json()
    assert result['thread_id']=='old' and 'pending' not in result['state']['followups'][0]
    assert len(result['state']['followups'][0]['messages'])==2
    assert len(client.get('/api/courses').json()['courses'])==1


def test_reasoning_files_are_private_exercise_bound_and_sent_with_attempt(client,monkeypatch):
    signup(client)
    cid=client.post('/api/courses',json={'profile':{'goal':'Python'}}).json()['id'];path=f'/api/courses/{cid}'
    assert client.post(path+'/reasoning',files={'file':('work.md',b'x = 3','text/markdown')}).status_code==400
    for action in ['approve','teach','practice']:assert client.post(path,json={'action':action}).status_code==200
    uploaded=client.post(path+'/reasoning',files={'file':('work.md',b'x = 3','text/markdown')})
    assert uploaded.status_code==200,uploaded.text
    rid=uploaded.json()['file']['id'];assert uploaded.json()['file']['text']=='x = 3'
    captured=[]
    def assess(state,answer,work):
        captured.append(work);return {'status':'correct','evidence':'Reasoning checked','possible_gap':'','diagnostic_question':''},{}
    monkeypatch.setattr(web_tutor,'assess',assess)
    result=client.post(path,json={'action':'attempt','answer':'3','work':'My reasoning','attachment_ids':[rid]})
    assert result.status_code==200,result.text
    assert 'x = 3' in captured[0] and 'My reasoning' in captured[0]
    client.post(path,json={'action':'practice'})
    owner=web_store.account(client.cookies.get('learning_login'))['id'];state,revision=web_store.course(owner,cid)
    state['question']['prompt']='A different exercise';web_store.save_course(owner,cid,state,revision)
    assert client.post(path,json={'action':'attempt','answer':'3','attachment_ids':[rid]}).status_code==400
    client.post('/api/logout');signup(client,'other-file-owner@example.test')
    assert client.post(path+'/reasoning',files={'file':('work.md',b'x = 3','text/markdown')}).status_code==400
    assert client.post(path,json={'action':'attempt','answer':'3','attachment_ids':[rid]}).status_code==400


def test_auto_library_notes_are_deduplicated_on_teach_and_reopen(client):
    signup(client)
    cid=client.post('/api/courses',json={'profile':{'goal':'Python'}}).json()['id'];path=f'/api/courses/{cid}'
    client.post(path,json={'action':'approve'});client.post(path,json={'action':'teach'})
    for _ in range(3):client.get(path)
    notes=client.get('/api/resources').json()['resources']
    assert len(notes)==1 and notes[0]['kind']=='note' and notes[0]['course_id']==cid
    client.post(path,json={'action':'teach'})
    assert len(client.get('/api/resources').json()['resources'])==1


def test_all_checkpoint_formats_omit_questions_and_render_clean_text():
    import checkpoints
    from pypdf import PdfReader
    state={'profile':{'goal':'Algebra'},'curriculum':{'title':'Algebra','objectives':[]},
           'lesson':{'explanation':r'A **variable** is $x$. {execution}Use $\frac{3}{4}$ and $x^{2}$.','worked_example':['Set $x = 4$.']},
           'question':{'prompt':'SECRET PRACTICE QUESTION'},'attempts':[{'question':'SECRET ATTEMPT QUESTION'}],
           'followups':[{'question':'SECRET CHAT QUESTION','answer':'SECRET CHAT ANSWER'}]}
    texts=[]
    for format in ['md','docx','pdf']:
        content,_=checkpoints.export(state,format)
        if format=='md':text=content.decode()
        elif format=='docx':text='\n'.join(p.text for p in Document(io.BytesIO(content)).paragraphs)
        else:text='\n'.join(p.extract_text() for p in PdfReader(io.BytesIO(content)).pages)
        texts.append(text)
        assert 'SECRET' not in text and '**' not in text and '$' not in text and '{execution}' not in text and r'\frac' not in text
        assert 'variable is x' in text and '(3)/(4)' in text and 'x²' in text
    assert 'Saved state' not in texts[0]
    assert checkpoints.readable_text('Use `x**2 + y**2` in code.')=='Use x**2 + y**2 in code.'


def test_short_docx_pdf_and_image_reasoning_extraction(monkeypatch):
    import reasoning_files
    from PIL import Image
    from reportlab.pdfgen.canvas import Canvas
    doc=Document();doc.add_paragraph('x = 3');buffer=io.BytesIO();doc.save(buffer)
    assert reasoning_files.extract(buffer.getvalue(),'work.docx')['text']=='x = 3'
    buffer=io.BytesIO();pdf=Canvas(buffer);pdf.drawString(30,700,'x = 3');pdf.save()
    assert 'x = 3' in reasoning_files.extract(buffer.getvalue(),'work.pdf')['text']
    image=Image.new('RGB',(50,50),'white');buffer=io.BytesIO();image.save(buffer,format='PNG')
    class Result:returncode=0;stdout=b'x = 3'
    monkeypatch.setattr(reasoning_files.subprocess,'run',lambda *a,**k:Result())
    monkeypatch.setattr(reasoning_files.shutil,'which',lambda name:'powershell.exe')
    assert reasoning_files.extract(buffer.getvalue(),'work.png')['text']=='x = 3'
    with pytest.raises(ValueError):reasoning_files.extract(b'not an image','bad.png')
    with pytest.raises(ValueError):reasoning_files.extract(b'bad','bad.exe')


def test_attachment_assessment_bounds_source_without_dropping_reasoning(monkeypatch):
    captured={}
    def generate(schema,purpose,data,instructions):
        captured.update(data=data)
        return {'status':'correct','evidence':'Reviewed work','possible_gap':'','diagnostic_question':''},{}
    monkeypatch.setattr(web_tutor.tutor,'generate',generate)
    state={'source':{'id':'reference','text':'文'*9000},'profile':{'goal':'Algebra'},'objective_index':0,'curriculum':{'objectives':[{'id':'o'}]},'question':{'prompt':'Solve x = 3','answer':'3','solution':'Read the value'}}
    work='Attached learner reasoning: '+('x = 3. '*500)
    web_tutor.assess(state,'3',work)
    assert captured['data']['work']==work and captured['data']['source_excerpt_only']
    assert len(json.dumps(captured['data'],ensure_ascii=False).encode())<18000
    assert state['source']['text']=='文'*9000


def test_similar_main_goals_offer_owned_learning_without_model_call(client,monkeypatch):
    signup(client)
    first=client.post('/api/courses',json={'profile':{'goal':'I want to learn C++'}}).json()['id']
    def fail(*args):raise AssertionError('Similar-goal check must not invoke inference')
    original=web_tutor.propose;monkeypatch.setattr(web_tutor,'propose',fail)
    result=client.post('/api/courses',json={'profile':{'goal':'Explain C++ to me?'}})
    assert result.status_code==200 and result.json()['similar_learning']['id']==first
    assert len(client.get('/api/courses').json()['courses'])==1
    monkeypatch.setattr(web_tutor,'propose',original)
    separate=client.post('/api/courses',json={'profile':{'goal':'Explain C++ to me'},'similar_choice':'new'}).json()
    assert separate['id']!=first
    client.post('/api/logout');signup(client,'other-course-owner@example.test')
    result=client.post('/api/courses',json={'profile':{'goal':'Explain C++ to me'}}).json()
    assert 'similar_learning' not in result
    from question_matching import goal_key
    assert goal_key('Learn C++')!=goal_key('Learn C')!=goal_key('Learn C#')
    assert goal_key('Learn Python loops')!=goal_key('Learn Python classes')


def test_saved_lessons_and_tests_survive_advancement_without_rewinding(client,monkeypatch):
    signup(client)
    cid=client.post('/api/courses',json={'profile':{'goal':'Python'}}).json()['id'];path=f'/api/courses/{cid}'
    owner=web_store.account(client.cookies.get('learning_login'))['id']
    state,revision=web_store.course(owner,cid)
    state['curriculum']['objectives']=[{'id':f'o{i}','title':f'Lesson {i}','outcome':'Learn','prerequisites':[],'source_ids':['general'],'minutes':5} for i in range(4)]
    web_store.save_course(owner,cid,state,revision)
    monkeypatch.setattr(web_tutor,'assess',lambda *a:({'status':'correct','evidence':'Checked','possible_gap':'','diagnostic_question':''},{}))
    client.post(path,json={'action':'approve'})
    for index in range(3):
        client.post(path,json={'action':'teach'})
        current=client.post(path,json={'action':'practice'}).json()['state']
        assert 'answer' not in current['lesson_history'][-1]['practices'][0]['question']
        client.post(path,json={'action':'attempt','answer':'3','work':'Assign 3'})
        response=client.post(path,json={'action':'next'})
        assert response.status_code==200,response.text
    state=client.get(path).json()['state']
    assert state['objective_index']==3 and state['stage']=='ready'
    assert len(state['lesson_history'])==3 and 'lesson' not in state
    first=state['lesson_history'][0]
    assert first['title']=='Lesson 0' and first['lesson']['explanation']
    assert first['practices'][0]['attempt']['answer']=='3' and first['practices'][0]['assessment']['status']=='correct'
    assert client.get(path).json()['state']['objective_index']==3


def test_repeated_tests_have_distinct_history_and_pending_solution_stays_private(client,monkeypatch):
    signup(client);cid=client.post('/api/courses',json={'profile':{'goal':'Python'}}).json()['id'];path=f'/api/courses/{cid}'
    monkeypatch.setattr(web_tutor,'assess',lambda *a:({'status':'correct','evidence':'Checked','possible_gap':'','diagnostic_question':''},{}))
    for action in ['approve','teach','practice']:client.post(path,json={'action':action})
    client.post(path,json={'action':'attempt','answer':'3'});result=client.post(path,json={'action':'practice'}).json()['state']
    tests=result['lesson_history'][0]['practices']
    assert len(tests)==2 and tests[0]['id']!=tests[1]['id']
    assert tests[0]['question']['solution'] and 'solution' not in tests[1]['question'] and 'assessment' not in tests[1]


def test_language_resolution_and_mixed_plan_warning():
    import language_rules as language
    assert language.resolve({'goal':'I want to learn C++','language':'automatic'})=='English'
    assert language.resolve({'goal':'我想学习代数','language':'automatic'})=='Mandarin Chinese'
    assert language.resolve({'goal':'saya ingin belajar Python'})=='Indonesian'
    assert language.resolve({'goal':'I want to learn C++','language':'French'})=='French'
    assert language.resolve({'goal':'I want to learn C++'},'我想再看一个例子')=='Mandarin Chinese'
    assert language.plan_mismatch({'goal':'I want to learn C++'},{'objectives':[{'outcome':'Identifier des variables, des types de données fondamentaux et effectuer des opérations arithmétiques simples.'}]})


def test_model_transport_receives_concrete_language_for_every_text_field(monkeypatch):
    from types import SimpleNamespace
    captured={}
    def request(settings,messages,**kwargs):
        captured.update(messages=messages)
        return SimpleNamespace(text=json.dumps({'answer':'A variable stores a value.'}),elapsed_seconds=0,usage={},response_model='fixture')
    monkeypatch.setattr(web_tutor.tutor,'request_model',request);monkeypatch.setattr(web_tutor.tutor,'load_settings',lambda:None)
    web_tutor.tutor.generate(web_tutor.Followup,'language_test',{'profile':{'goal':'I want to learn C++','language':'automatic'}},'Explain variables.')
    payload=json.loads(captured['messages'][1]['content'])
    assert payload['profile']['language']=='English' and payload['response_language']=='English'
    assert 'REQUIRED OUTPUT LANGUAGE: English' in captured['messages'][0]['content']


def test_history_keeps_pending_diagnostic_solution_private(client):
    signup(client);cid=client.post('/api/courses',json={'profile':{'goal':'Python'}}).json()['id'];path=f'/api/courses/{cid}'
    for action in ['approve','teach','practice']:client.post(path,json={'action':action})
    result=client.post(path,json={'action':'attempt','answer':'4'}).json()['state']
    assert result['stage']=='diagnostic'
    assert 'solution' not in result['lesson_history'][0]['practices'][0]['question']
    result=client.post(path,json={'action':'diagnostic','text':'I added one instead of assigning.'}).json()['state']
    assert result['lesson_history'][0]['practices'][0]['question']['solution']


def test_completed_checkpoint_keeps_saved_lessons_without_saved_tests():
    import checkpoints
    state={'profile':{'goal':'Python'},'curriculum':{'title':'Python','objectives':[]},'stage':'complete',
           'lesson_history':[{'title':'Variables','lesson':{'explanation':'Assignment binds names to values.','worked_example':['Set x = 3.']},'practices':[{'question':{'prompt':'PRIVATE TEST'}}]}]}
    text=checkpoints.export(state,'md')[0].decode()
    assert 'Assignment binds names' in text and 'Set x = 3' in text and 'PRIVATE TEST' not in text



def test_request_id_deduplication_and_structured_errors(client, monkeypatch):
    from llm import ModelRequestError
    signup(client)
    propose_calls = []
    orig_propose = web_tutor.propose
    def counting_propose(state, changes=""):
        propose_calls.append(1)
        return orig_propose(state, changes)
    monkeypatch.setattr(web_tutor, "propose", counting_propose)
    first = client.post("/api/courses", json={"profile": {"goal": "Learn algebra equations"}, "request_id": "create-req-1"}).json()
    second = client.post("/api/courses", json={"profile": {"goal": "Learn algebra equations"}, "request_id": "create-req-1"}).json()
    assert first["id"] == second["id"] and second.get("deduplicated") is True
    assert len(propose_calls) == 1
    assert "creation_request_id" not in second["state"]

    cid = first["id"]
    path = f"/api/courses/{cid}"
    client.post(path, json={"action": "approve", "request_id": "app-1"})
    teach_calls = []
    orig_teach = web_tutor.teach
    def counting_teach(state):
        teach_calls.append(1)
        return orig_teach(state)
    monkeypatch.setattr(web_tutor, "teach", counting_teach)
    t1 = client.post(path, json={"action": "teach", "request_id": "teach-req-1"}).json()
    t2 = client.post(path, json={"action": "teach", "request_id": "teach-req-1"}).json()
    assert len(teach_calls) == 1 and t2.get("deduplicated") is True
    assert "completed_requests" not in t1["state"]

    monkeypatch.setattr(web_tutor, "practice", lambda *a: (_ for _ in ()).throw(ModelRequestError("The request timed out.")))
    err_resp = client.post(path, json={"action": "practice", "request_id": "prac-timeout"})
    assert err_resp.status_code == 400
    err_json = err_resp.json()
    assert err_json["code"] == "timeout" and err_json["retryable"] is True


def test_episode_step_span_validation_assistance_tracking_and_observation_correction(client, monkeypatch):
    signup(client)
    cid = client.post("/api/courses", json={"profile": {"goal": "Solve linear equations"}}).json()["id"]
    path = f"/api/courses/{cid}"
    for action in ["approve", "teach", "practice"]:
        assert client.post(path, json={"action": action}).status_code == 200

    # 1. Wrong attempt with multi-line equation work triggers validated operator span and corrective_feedback
    work_lines = "x + 3 = 8\nx = 8 + 3\nx = 11"
    attempt_res = client.post(path, json={"action": "attempt", "answer": "x = 11", "work": work_lines}).json()["state"]
    assert attempt_res["stage"] == "diagnostic"
    episode = attempt_res["episode"]
    assert episode["micro_step"] == 1
    assert episode["work_sheet"]["has_flagged_step"] is True
    flagged = episode["work_sheet"]["flagged_span"]
    assert flagged["line_index"] == 1 and flagged["quoted_text"] == "+"
    assert episode["work_sheet"]["lines"][1]["text"][flagged["start"]:flagged["end"]] == "+"
    assert episode["outcome"]["corrective_feedback"] is True

    # 2. Learner can correct the sanitized observation; wrong observation_id or another user is rejected
    obs_id = episode["observation"]["id"]
    bad_obs = client.post(path, json={"action": "correct_observation", "observation_id": "wrong-id", "text": "Working with negative numbers"})
    assert bad_obs.status_code == 400
    good_obs = client.post(path, json={"action": "correct_observation", "observation_id": obs_id, "text": "Working with negative numbers"}).json()["state"]
    assert good_obs["episode"]["observation"]["corrected"] is True
    assert "Working with negative numbers" in good_obs["episode"]["observation"]["text"]
    assert "learning_signals" not in good_obs

    # 3. Diagnostic repair moves to micro_step 2 and records worked_solution_shown
    rep_state = client.post(path, json={"action": "diagnostic", "text": "Subtract 3 from both sides so x = 5"}).json()["state"]
    assert rep_state["stage"] == "repair"
    assert rep_state["episode"]["micro_step"] == 2
    assert rep_state["episode"]["outcome"]["worked_solution_shown"] is True

    # 4. Starting fresh transfer practice moves to micro_step 3 with clean assistance for the new instance
    prac2 = client.post(path, json={"action": "practice"}).json()["state"]
    assert prac2["stage"] == "attempt"
    assert prac2["episode"]["micro_step"] == 3
    assert prac2["episode"]["outcome"]["hint_opened"] is False

    # 5. Solving the transfer problem without opening a hint yields 'Solved independently'
    monkeypatch.setattr(web_tutor, "assess", lambda *a: ({"status": "correct", "evidence": "Subtracted 3 on both sides.", "possible_gap": "", "diagnostic_question": "Checked."}, {}))
    solved_ind = client.post(path, json={"action": "attempt", "answer": "3", "work": "x = 3"}).json()["state"]
    assert solved_ind["stage"] == "feedback"
    assert solved_ind["episode"]["outcome"]["category"] == "correct_without_recorded_help"
    assert solved_ind["episode"]["outcome"]["label"] == "Solved independently"
    assert solved_ind["episode"]["outcome"]["caveat"] == "One attempt is not proof of mastery."

    # 6. Opening a hint on another practice problem records hint_opened and changes outcome to 'Solved with a hint'
    client.post(path, json={"action": "practice"})
    hint_state = client.post(path, json={"action": "hint"}).json()["state"]
    assert hint_state["episode"]["hint"]
    assert hint_state["episode"]["outcome"]["hint_opened"] is True
    solved_hint = client.post(path, json={"action": "attempt", "answer": "3", "work": "x = 3"}).json()["state"]
    assert solved_hint["episode"]["outcome"]["category"] == "correct_with_help"
    assert solved_hint["episode"]["outcome"]["label"] == "Solved with a hint"


def test_google_callback_url_resolution(monkeypatch):
    from starlette.requests import Request

    # 1. Default local fallback
    scope_local = {
        "type": "http",
        "method": "GET",
        "path": "/auth/google",
        "headers": [(b"host", b"127.0.0.1:8780")],
        "server": ("127.0.0.1", 8780),
        "scheme": "http",
        "query_string": b"",
        "app": web_server.app,
    }
    assert web_server.google_callback_url(Request(scope_local)) == "http://127.0.0.1:8780/auth/google/callback"

    # 2. Reverse proxy / VS Code tunnel / Hugging Face Spaces forwarded headers
    scope_forwarded = {
        "type": "http",
        "method": "GET",
        "path": "/auth/google",
        "headers": [
            (b"host", b"localhost:8780"),
            (b"x-forwarded-proto", b"https"),
            (b"x-forwarded-host", b"tunnel-test.app.github.dev"),
        ],
        "server": ("127.0.0.1", 8780),
        "scheme": "http",
        "query_string": b"",
        "app": web_server.app,
    }
    assert web_server.google_callback_url(Request(scope_forwarded)) == "https://tunnel-test.app.github.dev/auth/google/callback"

    # 3. Explicit APP_BASE_URL environment variable
    monkeypatch.setenv("APP_BASE_URL", "https://my-space.hf.space")
    assert web_server.google_callback_url(Request(scope_forwarded)) == "https://my-space.hf.space/auth/google/callback"

    # 4. Explicit GOOGLE_REDIRECT_URI environment variable (highest precedence)
    monkeypatch.setenv("GOOGLE_REDIRECT_URI", "https://custom.example.org/oauth/callback")
    assert web_server.google_callback_url(Request(scope_forwarded)) == "https://custom.example.org/oauth/callback"

