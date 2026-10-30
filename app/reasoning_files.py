import io
import os
from pathlib import Path
import shutil
import subprocess
import tempfile
import zipfile
from PIL import Image, UnidentifiedImageError
from docx import Document
from pypdf import PdfReader


def extract(data, name):
    suffix=Path(name).suffix.lower()
    if len(data)>2*1024*1024:
        raise ValueError('Reasoning files must be smaller than 2 MB.')
    if suffix=='.md':
        try:text=data.decode('utf-8-sig')
        except UnicodeDecodeError:raise ValueError('Save the Markdown file as UTF-8.') from None
    elif suffix=='.pdf':
        try:
            reader=PdfReader(io.BytesIO(data))
            if len(reader.pages)>12:raise ValueError('Use up to 12 reasoning pages.')
            text='\n'.join(page.extract_text() or '' for page in reader.pages)
        except ValueError:raise
        except Exception:raise ValueError('This PDF could not be read. Use a text PDF or upload a page as an image.') from None
    elif suffix=='.docx':
        try:
            with zipfile.ZipFile(io.BytesIO(data)) as archive:
                if sum(x.file_size for x in archive.infolist())>8*1024*1024:raise ValueError('This DOCX expands beyond the document limit.')
            doc=Document(io.BytesIO(data))
            text='\n'.join([p.text for p in doc.paragraphs]+[' | '.join(c.text for c in row.cells) for table in doc.tables for row in table.rows])
        except ValueError:raise
        except Exception:raise ValueError('This DOCX could not be read.') from None
    elif suffix in {'.png','.jpg','.jpeg','.webp'}:
        try:
            with Image.open(io.BytesIO(data)) as image:
                if image.width*image.height>12_000_000:
                    raise ValueError('Use an image smaller than 12 megapixels.')
                image.verify()
            with Image.open(io.BytesIO(data)) as image, tempfile.TemporaryDirectory(prefix='learning-ocr-') as directory:
                image=image.convert('RGB');image.thumbnail((2400,2400))
                path=Path(directory)/'reasoning.png';image.save(path)
                if os.name!='nt' or not shutil.which('powershell.exe'):
                    raise ValueError('Local image OCR is unavailable. Upload a text PDF, DOCX or Markdown file instead.')
                result=subprocess.run([shutil.which('powershell.exe'),'-NoProfile','-NonInteractive','-ExecutionPolicy','Bypass','-File',str(Path(__file__).with_name('ocr_image.ps1')),'-ImagePath',str(path.resolve())],capture_output=True,timeout=35)
                if result.returncode:
                    raise ValueError('The image could not be read by local OCR. Try a clearer image or a text document.')
                text=result.stdout.decode('utf-8-sig').strip()
        except (UnidentifiedImageError, OSError, Image.DecompressionBombError):
            raise ValueError('This image could not be opened.') from None
        except subprocess.TimeoutExpired:
            raise ValueError('Local image reading timed out. Try a smaller image or paste the reasoning text.') from None
    else:
        raise ValueError('Choose Markdown, PDF, DOCX, PNG, JPEG or WebP.')
    if not text.strip():
        raise ValueError('No text was found. Try a clearer image or paste your reasoning.')
    if len(text.encode('utf-8'))>5000:
        raise ValueError('Use a shorter reasoning excerpt, up to 5 KB of extracted text.')
    return {'name':Path(name).name[:160],'text':text,'kind':'image' if suffix in {'.png','.jpg','.jpeg','.webp'} else suffix[1:]}
