"""Bounded reference extraction; fetched documents remain untrusted model input."""
import http.client
import ipaddress
import io
import os
import shutil
import socket
import ssl
import subprocess
import tempfile
import zipfile
from pathlib import Path
from urllib.parse import urlsplit, urljoin

from bs4 import BeautifulSoup
from docx import Document
import requests

from content import pdf_source

MAX_TEXT = 9000
VIDEO_SUFFIXES = {".mp4", ".webm", ".mov", ".mkv", ".mp3", ".wav", ".m4a"}


def source(text, title, origin, kind="text"):
    text = text.strip()
    if len(text) < 20:
        raise ValueError("The reference has too little readable text. Supply a transcript or text excerpt.")
    if len(text) > MAX_TEXT:
        raise ValueError("Use an excerpt of at most 9,000 characters.")
    return {"id": kind, "title": title[:160], "origin": origin[:1000], "text": text}


def document(data, name):
    suffix = Path(name).suffix.lower()
    if suffix in VIDEO_SUFFIXES:
        return transcribe_video(data, name)
    if len(data) > 2 * 1024 * 1024:
        raise ValueError("Documents must be smaller than 2 MB.")
    if suffix == ".pdf":
        return pdf_source(data, Path(name).name)
    if suffix == ".docx":
        try:
            with zipfile.ZipFile(io.BytesIO(data)) as archive:
                if sum(x.file_size for x in archive.infolist()) > 8 * 1024 * 1024:
                    raise ValueError("This DOCX expands beyond the document limit.")
            doc = Document(io.BytesIO(data))
            text = "\n".join([p.text for p in doc.paragraphs] +
                             [" | ".join(c.text for c in row.cells) for t in doc.tables for row in t.rows])
        except ValueError:
            raise
        except Exception:
            raise ValueError("This DOCX could not be read.") from None
        return source(text, Path(name).name, "Uploaded DOCX", "docx")
    if suffix in {".txt", ".md", ".srt", ".vtt"}:
        try:
            return source(data.decode("utf-8-sig"), Path(name).name, "Uploaded text or transcript", "text")
        except UnicodeDecodeError:
            raise ValueError("Save the text file as UTF-8.") from None
    raise ValueError("Supported: PDF, DOCX, TXT, Markdown, SRT/VTT, and configured audio/video uploads.")


def public_target(url):
    parsed = urlsplit(url)
    if parsed.scheme not in {"http", "https"} or not parsed.hostname or parsed.username or parsed.password:
        raise ValueError("Use a public HTTP or HTTPS link without credentials.")
    port = parsed.port or (443 if parsed.scheme == "https" else 80)
    if port not in {80, 443}:
        raise ValueError("Only standard public website ports are supported.")
    try:
        addresses = socket.getaddrinfo(parsed.hostname, port, type=socket.SOCK_STREAM)
    except OSError:
        raise ValueError("This website could not be resolved.") from None
    ips = [a[4][0] for a in addresses]
    if not ips or any(not ipaddress.ip_address(ip).is_global for ip in ips):
        raise ValueError("Private, local, and reserved network addresses cannot be read.")
    return parsed, port, ips[0]


class PinnedHTTPS(http.client.HTTPSConnection):
    def __init__(self, hostname, address, port):
        super().__init__(hostname, port, timeout=10, context=ssl.create_default_context())
        self.address = address

    def connect(self):
        raw = socket.create_connection((self.address, self.port), self.timeout)
        self.sock = self._context.wrap_socket(raw, server_hostname=self.host)


def read_url(url):
    url = url.strip()
    for _ in range(4):
        parsed, port, address = public_target(url)
        connection = PinnedHTTPS(parsed.hostname, address, port) if parsed.scheme == "https" else http.client.HTTPConnection(address, port, timeout=10)
        try:
            path = parsed.path or "/"
            if parsed.query:
                path += "?" + parsed.query
            connection.request("GET", path, headers={"Host": parsed.hostname, "User-Agent": "PersonalLearning/1.0", "Accept-Encoding": "identity"})
            response = connection.getresponse()
            if response.status in {301, 302, 303, 307, 308}:
                location = response.getheader("Location")
                if not location:
                    raise ValueError("This link redirects without a destination.")
                url = urljoin(url, location)
                continue
            if response.status != 200:
                raise ValueError("The page is unavailable or requires login. Upload or paste the relevant text.")
            mime = response.getheader("Content-Type", "").lower()
            if not any(t in mime for t in ("text/html", "text/plain", "application/pdf")):
                raise ValueError("This link is not an HTML page, text, or PDF. Upload a supported file instead.")
            data = response.read(2 * 1024 * 1024 + 1)
            if len(data) > 2 * 1024 * 1024:
                raise ValueError("This page exceeds the 2 MB download limit.")
            if "application/pdf" in mime:
                result = pdf_source(data, Path(parsed.path).name or "Linked PDF")
                result.update(origin=url)
                return result
            if "text/plain" in mime:
                return source(data.decode("utf-8", errors="replace"), parsed.hostname, url, "url")
            soup = BeautifulSoup(data, "html.parser")
            title = soup.title.get_text(" ", strip=True) if soup.title else parsed.hostname
            for node in soup(["script", "style", "nav", "footer", "header", "form"]):
                node.decompose()
            main = soup.find("main") or soup.find("article") or soup
            text = main.get_text(" ", strip=True)
            # The preview tells the learner when only the opening excerpt is retained.
            result = source(text[:MAX_TEXT], title, url, "url")
            result["excerpted"] = len(text) > MAX_TEXT
            return result
        except (OSError, http.client.HTTPException):
            raise ValueError("Could not read this public page. Upload or paste its content instead.") from None
        finally:
            connection.close()
    raise ValueError("This link redirects too many times.")


def transcription_available():
    return bool(shutil.which("ffmpeg") and os.getenv("TRANSCRIPTION_BASE_URL") and os.getenv("TRANSCRIPTION_MODEL") and os.getenv("TRANSCRIPTION_API_KEY") and os.getenv("TRANSCRIPTION_ALLOW_PAID", "false").lower() == "true")


def transcribe_video(data, name):
    if len(data) > 25 * 1024 * 1024:
        raise ValueError("Use audio/video smaller than 25 MB and at most five minutes.")
    if not transcription_available():
        raise ValueError("Video transcription is not configured. Upload its SRT/VTT/TXT transcript now, or configure the transcription service described in docs/INTEGRATION.md.")
    with tempfile.TemporaryDirectory(prefix="learning-video-") as directory:
        incoming = Path(directory) / ("input" + Path(name).suffix.lower())
        audio = Path(directory) / "audio.wav"
        incoming.write_bytes(data)
        try:
            formats = {".mp4":"mov", ".mov":"mov", ".m4a":"mov", ".webm":"matroska", ".mkv":"matroska", ".mp3":"mp3", ".wav":"wav"}
            subprocess.run([shutil.which("ffmpeg"), "-nostdin", "-v", "error", "-protocol_whitelist", "file,pipe", "-f", formats[Path(name).suffix.lower()], "-i", str(incoming), "-t", "301", "-vn", "-ac", "1", "-ar", "16000", str(audio)], check=True, capture_output=True, timeout=90)
            import wave
            with wave.open(str(audio)) as wav:
                seconds = wav.getnframes() / wav.getframerate()
            if seconds > 300:
                raise ValueError("Use a video no longer than five minutes.")
            base = os.getenv("TRANSCRIPTION_BASE_URL", "").rstrip("/")
            if not base.startswith("https://"):
                raise ValueError("Transcription requires an HTTPS service endpoint.")
            with audio.open("rb") as file:
                response = requests.post(base + "/audio/transcriptions", headers={"Authorization": "Bearer " + os.getenv("TRANSCRIPTION_API_KEY", "")}, files={"file": ("audio.wav", file, "audio/wav")}, data={"model": os.getenv("TRANSCRIPTION_MODEL"), "response_format": "json"}, timeout=120)
            if response.status_code != 200:
                raise ValueError("The transcription service could not process this file. Upload a transcript instead.")
            result = source(response.json().get("text", ""), Path(name).name, "Audio transcript; visual slides and diagrams are not analysed", "transcript")
            return result
        except (subprocess.SubprocessError, requests.RequestException, OSError):
            raise ValueError("Could not transcribe this video. Upload its transcript instead.") from None
